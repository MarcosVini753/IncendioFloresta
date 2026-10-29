"""Gera PDF e Word do estudo; não altera produtos científicos ou o frontend."""
from __future__ import annotations

import html
import json
import re
from decimal import Decimal as D
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle

from calcular_cenarios import calculate

HERE = Path(__file__).resolve().parent
STEM = 'estudo-servidor-risco-v4-2026-10-26'
GREEN = '#173F35'
ORANGE = '#B34F25'
LIGHT = '#EAF2EE'


def number(value, decimals=2):
    return f'{value:,.{decimals}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def money(value):
    return 'R$ ' + number(value)


def mdtable(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join('---' for _ in headers) + ' |'] + ['| ' + ' | '.join(str(v) for v in row) + ' |' for row in rows])


def replacements(raw, result):
    rows = {r['id']: r for r in result['rows']}
    cpu, gpu = result['cap_cpu'], result['cap_gpu']
    hardware = raw['hardware']
    pc = hardware[:7]
    facilities = hardware[7:]
    gpu_line = raw['gpu_local']

    def parts_table(items):
        return mdtable(['Item', 'Valor de referência', 'Natureza / fonte'], [[r['item'], money(D(str(r['valor']))), r['tipo'] + (' [' + r['fonte'] + ']' if r['fonte'] else '')] for r in items])

    sources = []
    for ref in raw['fontes']:
        sources.append(f"**[{ref['id']}] {ref['titulo']}.** {ref['evidencia']} [Abrir fonte]({ref['url']}).")
    p = raw['premissas']
    return {
        'resumo_table': mdtable(['Alternativa', 'Referência de custo', 'Leitura correta'], [
            ['Nuvem + B', money(rows['nuvem_b']['anual']) + '/ano', 'Infraestrutura do ciclo anual e publicação; trabalho técnico separado'],
            ['Nuvem + C', money(rows['nuvem_c']['anual']) + '/ano', 'Mesma ciência + VPS de mapas como entrada para piloto'],
            ['Próprio sem GPU', money(cpu) + ' iniciais', 'Inclui instalação e reserva; operação anual adicional'],
            ['Próprio com GPU', money(gpu) + ' iniciais', 'Configuração cotada com 64 GB, SSD de 2 TB e GPU de 16 GB']]),
        'm': number(result['m'], 0),
        'node_visits': number(result['node_visits'], 0),
        'throughput_table': mdtable(['Taxa efetiva hipotética', 'Conta aplicada', 'Só árvores'], [[number(rate, 0) + ' visitas/s', '12,89 trilhões / taxa / 3.600', number(result['node_visits'] / rate / 3600) + ' h'] for rate in [100000000, 300000000, 600000000, 1000000000]]),
        'equipment_table': parts_table(pc + [gpu_line]),
        'facilities_table': parts_table(facilities),
        'capex_table': mdtable(['Compra própria', 'Subtotal', 'Reserva de 15%', 'Total inicial'], [[
            'CPU, sem GPU dedicada', money(result['equipment_subtotal']), money(result['equipment_subtotal'] * D('.15')), money(cpu)], [
            'GPU, base + RTX 5060 Ti', money(result['equipment_subtotal'] + D(str(gpu_line['valor']))), money((result['equipment_subtotal'] + D(str(gpu_line['valor']))) * D('.15')), money(gpu)]]),
        'cloud_job': money(result['cloud_job']),
        'cloud_b_annual': money(rows['nuvem_b']['anual']),
        'cloud_c_annual': money(rows['nuvem_c']['anual']),
        'c_vps': money(result['c_vps']),
        'tco_infra_table': mdtable(['Cenário', 'Anual', '1 ano', '3 anos', '5 anos'], [[r['nome'], money(r['anual']), *[money(r['tco_infra'][str(y)]) for y in [1, 3, 5]]] for r in result['rows']]),
        'tco_work_table': mdtable(['Cenário', 'Implantação técnica', 'Trabalho / ano', 'Total 3 anos', 'Total 5 anos'], [[r['nome'], money(r['implantacao_tecnica']), money(r['operacao_tecnica_ano']), money(r['tco_com_trabalho']['3']), money(r['tco_com_trabalho']['5'])] for r in result['rows']]),
        'traffic_table': mdtable(['Visitas / mês', 'B: transferência', 'C: transferência', 'C: CPU estimada'], [[number(r['visitas'], 0), number(r['b_gb'], 0) + ' GB/mês', number(r['c_gb'], 0) + ' GB/mês', number(r['c_cpu_h_mes']) + ' h/mês'] for r in result['traffic']]),
        'sensitivity_table': mdtable(['Reserva de máquina / ano', 'Infra anual B', 'Infra anual C'], [[str(r['horas']) + ' h', money(r['b']), money(r['c'])] for r in result['sensitivity']]),
        'always_on': money(result['gpu_24x7_brl_ano_sem_disco']),
        'fx_low_b': money((rows['nuvem_b']['anual'] - D(str(p['dominio_brl_ano']))) * D('.8') + D(str(p['dominio_brl_ano']))),
        'fx_high_b': money((rows['nuvem_b']['anual'] - D(str(p['dominio_brl_ano']))) * D('1.2') + D(str(p['dominio_brl_ano']))),
        'break_even': number(min(result['break_even_cloud_hours'].values()), 0) + '–' + number(max(result['break_even_cloud_hours'].values()), 0),
        'source_list': '\n\n'.join(sources),
    }


def blocks(text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == '<!-- pagebreak -->':
            yield 'pagebreak', None
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v or '-') for v in cells):
                    rows.append(cells)
                i += 1
            yield 'table', rows
            continue
        elif line.startswith('# '):
            yield 'title', line[2:]
        elif line.startswith('## '):
            yield 'h2', line[3:]
        elif line.startswith('### '):
            yield 'h3', line[4:]
        elif line.startswith('> '):
            yield 'quote', line[2:]
        elif line.startswith('- '):
            yield 'bullet', line[2:]
        elif re.match(r'^\d+\. ', line):
            yield 'number', line
        else:
            paragraph = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r'^(#|\||>|- |\d+\. |<!--)', lines[i + 1].strip()):
                i += 1
                paragraph.append(lines[i].strip())
            yield 'p', ' '.join(paragraph)
        i += 1


def rich(text):
    text = html.escape(text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<link href="\2" color="#245F88">\1</link>', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
    return re.sub(r'`([^`]+)`', r'<font name="DVMono" size="8">\1</font>', text)


def table_widths(rows, width):
    n = len(rows[0])
    if rows[0][0] == 'Perfil possível':
        return [width * v for v in [.22, .40, .38]]
    weights = {2: [0.48, 0.52], 3: [0.46, 0.25, 0.29], 4: [0.34, 0.22, 0.22, 0.22], 5: [0.25, 0.18, 0.19, 0.19, 0.19]}.get(n, [1 / n] * n)
    return [v * width for v in weights]


def pdf_document(content, path):
    fonts = Path('/usr/share/fonts/truetype/dejavu')
    for name, filename in [('DV', 'DejaVuSans.ttf'), ('DV-Bold', 'DejaVuSans-Bold.ttf'), ('DV-Italic', 'DejaVuSans-Oblique.ttf'), ('DV-BoldItalic', 'DejaVuSans-BoldOblique.ttf'), ('DVMono', 'DejaVuSansMono.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    pdfmetrics.registerFontFamily('DV', normal='DV', bold='DV-Bold', italic='DV-Italic', boldItalic='DV-BoldItalic')
    styles = {
        'p': ParagraphStyle('body', fontName='DV', fontSize=9.5, leading=14, textColor=colors.HexColor('#273A34'), spaceAfter=8),
        'title': ParagraphStyle('title', fontName='DV-Bold', fontSize=29, leading=36, textColor=colors.HexColor(GREEN), spaceBefore=26, spaceAfter=18, keepWithNext=True),
        'h2': ParagraphStyle('h2', fontName='DV-Bold', fontSize=17, leading=22, textColor=colors.HexColor(GREEN), spaceAfter=14, spaceBefore=4, keepWithNext=True),
        'h3': ParagraphStyle('h3', fontName='DV-Bold', fontSize=11.5, leading=16, textColor=colors.HexColor(ORANGE), spaceAfter=9, spaceBefore=10, keepWithNext=True),
        'quote': ParagraphStyle('quote', fontName='DV', fontSize=10, leading=15, textColor=colors.HexColor(GREEN), backColor=colors.HexColor(LIGHT), borderPadding=10, leftIndent=10, rightIndent=10, spaceBefore=5, spaceAfter=16),
        'table': ParagraphStyle('cell', fontName='DV', fontSize=7.9, leading=11.2, textColor=colors.HexColor('#24352F'), splitLongWords=True),
        'th': ParagraphStyle('th', fontName='DV-Bold', fontSize=7.9, leading=11.2, textColor=colors.white),
    }
    styles['bullet'] = ParagraphStyle('bullet', parent=styles['p'], leftIndent=12, firstLineIndent=-8, spaceAfter=5)
    styles['number'] = ParagraphStyle('number', parent=styles['p'], leftIndent=12, firstLineIndent=-12, spaceAfter=8)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=19 * mm, rightMargin=19 * mm, topMargin=22 * mm, bottomMargin=20 * mm,
                            title='Estudo de infraestrutura — Risco anual do Acre', author='Projeto IncendioFloresta', subject='Nuvem, servidor próprio, complexidade computacional e custo-benefício')
    story = []
    for kind, value in blocks(content):
        if kind == 'pagebreak':
            story.append(PageBreak())
        elif kind == 'table':
            cells = [[Paragraph(rich(v), styles['th' if i == 0 else 'table']) for v in row] for i, row in enumerate(value)]
            table = Table(cells, colWidths=table_widths(value, doc.width), repeatRows=1, hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(GREEN)), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#F0F5F2'), colors.white]), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7), ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5), ('LINEBELOW', (0, 0), (-1, 0), 0.6, colors.HexColor(GREEN)), ('LINEBELOW', (0, 1), (-1, -1), 0.3, colors.HexColor('#D2DFD7'))]))
            story.extend([table, Spacer(1, 12)])
        else:
            if kind == 'bullet':
                value = '• ' + value
            story.append(Paragraph(rich(value), styles[kind]))

    def decorations(canvas, page_doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor('#D2DFD7'))
        canvas.line(19 * mm, A4[1] - 15 * mm, A4[0] - 19 * mm, A4[1] - 15 * mm)
        canvas.setFont('DV-Bold', 7.2)
        canvas.setFillColor(colors.HexColor(GREEN))
        canvas.drawString(19 * mm, A4[1] - 12 * mm, 'INCENDIOFLORESTA  /  PLANEJAMENTO DE INFRAESTRUTURA')
        canvas.setFont('DV', 7)
        canvas.setFillColor(colors.HexColor('#65776F'))
        canvas.drawString(19 * mm, 12 * mm, '26/10/2026 · Estudo orçamentário · Valores e premissas sujeitos a revisão')
        canvas.drawRightString(A4[0] - 19 * mm, 12 * mm, str(page_doc.page))
        canvas.restoreState()
    doc.build(story, onFirstPage=decorations, onLaterPages=decorations)


def add_link(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    rel_id = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link.set(qn('r:id'), rel_id)
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '245F88')
    properties.append(color)
    run.append(properties)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def word_text(paragraph, text):
    pattern = r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for chunk in re.split(pattern, text):
        if not chunk:
            continue
        if chunk.startswith('**'):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith('`'):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif chunk.startswith('[') and '](' in chunk:
            match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', chunk)
            add_link(paragraph, *match.groups())
        else:
            paragraph.add_run(chunk)


def word_document(content, path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.9)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [('Title', 30), ('Heading 1', 18), ('Heading 2', 12)]:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string((ORANGE if name == 'Heading 2' else GREEN).lstrip('#'))
        st.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = 'INCENDIOFLORESTA  /  PLANEJAMENTO DE INFRAESTRUTURA'
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string('65776F')
    footer = section.footer.paragraphs[0]
    footer.add_run('26/10/2026 · Estudo orçamentário · Página ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.title = 'Servidor para o Risco anual do Acre — Nuvem ou equipamento próprio?'
    doc.core_properties.author = 'Projeto IncendioFloresta'
    for kind, value in blocks(content):
        if kind == 'pagebreak':
            doc.add_page_break()
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.autofit = False
            widths = table_widths(value, 17.2)
            for row_i, values in enumerate(value):
                row = table.add_row()
                props = row._tr.get_or_add_trPr()
                cant_split = OxmlElement('w:cantSplit')
                props.append(cant_split)
                if row_i == 0:
                    props.append(OxmlElement('w:tblHeader'))
                for col_i, text in enumerate(values):
                    cell = row.cells[col_i]
                    cell.width = Cm(widths[col_i])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    shading = OxmlElement('w:shd')
                    shading.set(qn('w:fill'), GREEN[1:] if row_i == 0 else ('F0F5F2' if row_i % 2 else 'FFFFFF'))
                    cell._tc.get_or_add_tcPr().append(shading)
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_before = Pt(4)
                    para.paragraph_format.space_after = Pt(4)
                    word_text(para, text)
                    for run in para.runs:
                        run.font.size = Pt(8.5)
                        if row_i == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        else:
            style = {'title': 'Title', 'h2': 'Heading 1', 'h3': 'Heading 2', 'bullet': 'List Bullet'}.get(kind, 'Normal')
            paragraph = doc.add_paragraph(style=style)
            if kind == 'quote':
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), LIGHT[1:])
                paragraph._p.get_or_add_pPr().append(shading)
                paragraph.paragraph_format.left_indent = Cm(.3)
                paragraph.paragraph_format.right_indent = Cm(.3)
            word_text(paragraph, value)
    doc.save(path)


def main():
    raw, result = calculate()
    # Verificações de coerência apenas do relatório; nenhum job científico.
    assert result['n'] == 172340796 and result['m'] == 43092640
    assert result['node_visits'] == 12891091540800
    rows = {r['id']: r for r in result['rows']}
    assert abs(rows['nuvem_b']['anual'] - D('877.762')) < D('.001')
    assert rows['proprio_gpu_b']['capex'] > rows['proprio_cpu_b']['capex']
    for row in result['rows']:
        for year in [1, 3, 5]:
            assert row['tco_com_trabalho'][str(year)] > row['tco_infra'][str(year)]
    content = (HERE / 'conteudo_estudo.md').read_text()
    for key, value in replacements(raw, result).items():
        content = content.replace('{{' + key + '}}', value)
    assert not re.search(r'\{\{[^}]+\}\}', content), 'Placeholder não substituído'
    (HERE / (STEM + '.md')).write_text(content)
    (HERE / 'resultados_calculo.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, default=lambda v: float(v) if isinstance(v, D) else v) + '\n')
    pdf_document(content, HERE / (STEM + '.pdf'))
    word_document(content, HERE / (STEM + '.docx'))
    for ext in ['pdf', 'docx', 'md']:
        path = HERE / (STEM + '.' + ext)
        print(path, path.stat().st_size, 'bytes')


if __name__ == '__main__':
    main()
